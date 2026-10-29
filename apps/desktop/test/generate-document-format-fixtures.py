#!/usr/bin/env python3
"""Generate synthetic inputs; --soffice is used only to export real DOC/PPT fixtures."""
import argparse
import hashlib
import io
import json
import shutil
import struct
import subprocess
import zipfile
import zlib
from pathlib import Path

import msoffcrypto
from docx import Document
from docx.enum.section import WD_SECTION
from docx.shared import Inches, Pt
from PIL import Image, ImageDraw
from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.util import Inches as PptInches, Pt as PptPt
from pypdf import PdfWriter
from pypdf.generic import DictionaryObject, NameObject, DecodedStreamObject


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('directory', type=Path)
    parser.add_argument('--soffice', default=shutil.which('soffice'))
    args = parser.parse_args()
    if not args.soffice:
        parser.error('A fixture-authoring LibreOffice CLI is required to produce real legacy DOC/PPT.')
    root = args.directory.resolve()
    root.mkdir(parents=True, exist_ok=False)
    image = Image.new('RGB', (320, 180), 'white')
    draw = ImageDraw.Draw(image)
    draw.rectangle((10, 10, 150, 170), fill='#275fa5')
    draw.ellipse((170, 20, 310, 160), fill='#d98436')
    image.save(root / 'valid.png')
    image.save(root / 'valid.jpg')
    image.save(root / 'valid.webp')
    image.save(root / 'valid.gif')
    image.save(root / 'unsupported.bmp')
    image.save(root / 'animated.gif', save_all=True, append_images=[Image.new('RGB', image.size, 'red')], duration=100, loop=0)

    doc = Document()
    doc.styles['Normal'].font.name = 'Arial'
    doc.styles['Normal'].font.size = Pt(12)
    doc.sections[0].header.paragraphs[0].text = 'SPARKX HEADER'
    doc.sections[0].footer.paragraphs[0].text = 'SPARKX FOOTER'
    doc.add_heading('SPARKX PORTRAIT PAGE', 0)
    doc.add_paragraph('中文兼容性测试。English text, 00123 and punctuation: () [] — €.')
    table = doc.add_table(rows=3, cols=3)
    table.style = 'Table Grid'
    for r in range(3):
        for c in range(3):
            table.cell(r, c).text = f'CELL {r+1}-{c+1}'
    table.cell(1, 0).merge(table.cell(1, 1)).text = 'MERGED CELL'
    doc.add_picture(str(root / 'valid.png'), width=Inches(3))
    section = doc.add_section(WD_SECTION.NEW_PAGE)
    section.page_width, section.page_height = Inches(11), Inches(8.5)
    doc.add_heading('SPARKX LANDSCAPE PAGE', 0)
    doc.add_paragraph('LANDSCAPE BODY / 中文横向页面')
    for text in ['First numbered paragraph', 'Second numbered paragraph']:
        doc.add_paragraph(text, style='List Number')
    doc.save(root / 'complex.docx')
    missing = Document()
    paragraph = missing.add_paragraph('MISSING FONT FALLBACK / 中文字体回退')
    paragraph.runs[0].font.name = 'SparkX Missing Font QA 20261029'
    missing.save(root / 'missing-font.docx')

    deck = Presentation()
    deck.slide_width, deck.slide_height = PptInches(13.333), PptInches(7.5)
    for n in range(1, 4):
        slide = deck.slides.add_slide(deck.slide_layouts[6])
        box = slide.shapes.add_textbox(PptInches(.5), PptInches(.4), PptInches(12), PptInches(1.3))
        box.text_frame.text = f'SPARKX SLIDE {n} / 中文幻灯片'
        box.text_frame.paragraphs[0].runs[0].font.size = PptPt(32)
        box.text_frame.paragraphs[0].runs[0].font.color.rgb = RGBColor(30, 60, 90)
        slide.shapes.add_picture(str(root / 'valid.png'), PptInches(.5), PptInches(2), width=PptInches(4))
        cells = slide.shapes.add_table(3, 2, PptInches(5), PptInches(2), PptInches(6), PptInches(3)).table
        for r in range(3):
            for c in range(2):
                cells.cell(r, c).text = f'S{n} R{r+1} C{c+1}'
    deck.save(root / 'complex.pptx')
    for extension, filter_name in [('doc', 'MS Word 97'), ('ppt', 'MS PowerPoint 97')]:
        source = root / ('complex.docx' if extension == 'doc' else 'complex.pptx')
        profile = root / f'authoring-profile-{extension}'
        result = subprocess.run([args.soffice, f'-env:UserInstallation={profile.as_uri()}', '--headless', '--convert-to',
                                 f'{extension}:{filter_name}', '--outdir', str(root), str(source)],
                                check=True, capture_output=True, text=True, timeout=60)
        target = root / f'complex.{extension}'
        assert target.read_bytes()[:8] == bytes.fromhex('d0cf11e0a1b11ae1'), (target, result.stdout, result.stderr)
        shutil.rmtree(profile, ignore_errors=True)

    for extension in ['docx', 'pptx']:
        raw = (root / f'complex.{extension}').read_bytes()
        target = root / f'encrypted.{extension}'
        with target.open('wb') as output:
            msoffcrypto.OfficeFile(io.BytesIO(raw)).encrypt('SparkX synthetic fixture password', output)
        encrypted = msoffcrypto.OfficeFile(io.BytesIO(target.read_bytes()))
        assert encrypted.is_encrypted()
        encrypted.load_key(password='SparkX synthetic fixture password')
        recovered = io.BytesIO()
        encrypted.decrypt(recovered)
        assert recovered.getvalue() == raw
        (root / f'corrupt.{extension}').write_bytes(raw[:64])
        with zipfile.ZipFile(io.BytesIO(raw)) as source, zipfile.ZipFile(root / f'embedded.{extension}', 'w') as output:
            for entry in source.infolist():
                output.writestr(entry, source.read(entry.filename))
            output.writestr(('word' if extension == 'docx' else 'ppt') + '/embeddings/unsupported.bin', b'not executable')
    (root / 'corrupt.doc').write_bytes(bytes.fromhex('d0cf11e0a1b11ae1') + bytes(504))
    (root / 'corrupt.ppt').write_bytes(bytes.fromhex('d0cf11e0a1b11ae1') + bytes(504))
    (root / 'wrong-signature.doc').write_text('This is text, not a DOC binary.', encoding='utf-8')
    (root / 'wrong-signature.ppt').write_text('This is text, not a PPT binary.', encoding='utf-8')

    writer = PdfWriter()
    for n, (width, height) in enumerate([(300, 420), (420, 300)], 1):
        page = writer.add_blank_page(width=width, height=height)
        font = DictionaryObject({NameObject('/Type'): NameObject('/Font'), NameObject('/Subtype'): NameObject('/Type1'), NameObject('/BaseFont'): NameObject('/Helvetica')})
        page[NameObject('/Resources')] = DictionaryObject({NameObject('/Font'): DictionaryObject({NameObject('/F1'): font})})
        stream = DecodedStreamObject()
        stream.set_data(f'BT /F1 20 Tf 20 {height-40} Td (SPARKX PDF PAGE {n}) Tj ET'.encode())
        page[NameObject('/Contents')] = writer._add_object(stream)
    with (root / 'valid.pdf').open('wb') as output:
        writer.write(output)
    writer.encrypt('SparkX synthetic fixture password', algorithm='AES-256')
    with (root / 'encrypted.pdf').open('wb') as output:
        writer.write(output)
    (root / 'corrupt.pdf').write_bytes(b'%PDF-1.7\nnot a document\n%%EOF')
    (root / 'corrupt.png').write_bytes((root / 'valid.png').read_bytes()[:40])
    oversized = bytearray((root / 'valid.png').read_bytes())
    oversized[16:24] = struct.pack('>II', 32768, 32768)
    oversized[29:33] = struct.pack('>I', zlib.crc32(oversized[12:29]))
    (root / 'oversized.png').write_bytes(oversized)
    (root / 'safe.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><rect width="100" height="100" fill="blue"/></svg>')
    (root / 'active.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><script>window.compromised=true</script></svg>')
    (root / 'external.svg').write_text('<svg xmlns="http://www.w3.org/2000/svg" width="100" height="100"><image href="https://example.invalid/private.png"/></svg>')
    (root / 'utf8.txt').write_text('中文 UTF-8 text\n<script>window.compromised=true</script>', encoding='utf-8')
    (root / 'boundary.txt').write_bytes(b'a' * (256*1024-1) + '中文 END'.encode())
    (root / 'binary.txt').write_bytes(b'text\0binary')
    (root / 'invalid-utf8.txt').write_bytes(bytes([0xff, 0xfe, 0x80]))
    (root / 'unsafe.md').write_text('# SYNTHETIC MARKDOWN\n\n<script>window.compromised=true</script>\n\n![remote](https://example.invalid/private.png)\n\n[link](https://example.invalid/)\n\n![escape](../outside.png)\n', encoding='utf-8')
    (root / 'valid.csv').write_text('code,note\n00123,"中文\nquoted line"\n=1+1,false', encoding='utf-8')
    (root / 'invalid.csv').write_bytes(b'"unclosed')
    (root / 'invalid-utf8.csv').write_bytes(bytes([0xff, 0xfe, 0x80]))
    manifest = {'schema_version': 1, 'synthetic_only': True,
                'generator': 'python-docx/python-pptx/Pillow/pypdf/msoffcrypto; LibreOffice CLI legacy export',
                'files': {p.name: {'bytes': p.stat().st_size, 'sha256': hashlib.sha256(p.read_bytes()).hexdigest()}
                          for p in sorted(root.iterdir()) if p.is_file()}}
    (root / 'manifest.json').write_text(json.dumps(manifest, indent=2) + '\n')
    print(json.dumps({'directory': str(root), 'fixtures': len(manifest['files'])}))


if __name__ == '__main__':
    main()
