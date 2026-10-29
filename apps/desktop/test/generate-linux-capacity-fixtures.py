"""Generate synthetic Linux capacity fixtures; requires Pillow, python-docx and python-pptx."""
from pathlib import Path
import hashlib, json, random, sys
from PIL import Image, ImageDraw
from docx import Document
from docx.shared import Inches as WordInches
from pptx import Presentation
from pptx.util import Inches, Pt
out = Path(sys.argv[1]); out.mkdir(parents=True, exist_ok=True)
images=[]
for n in range(8):
    rng=random.Random(20261029+n)
    img=Image.frombytes('RGB',(640,400),rng.randbytes(640*400*3))
    draw=ImageDraw.Draw(img);draw.rectangle((0,0,640,70),fill='white');draw.text((15,20),f'SYNTHETIC CAPACITY IMAGE {n+1}',fill='black')
    p=out/f'image-{n}.png';img.save(p);images.append(p)
doc=Document()
for n in range(40):
    if n: doc.add_page_break()
    doc.add_heading(f'SYNTHETIC CAPACITY PAGE {n+1:02d}',level=1)
    doc.add_paragraph('Synthetic English text and Chinese text: 中文文档兼容性验收。 Every paragraph belongs to an owned qualification fixture.')
    table=doc.add_table(rows=6,cols=3);table.style='Table Grid'
    for i,row in enumerate(table.rows):
        for j,cell in enumerate(row.cells):cell.text=f'Page {n+1}, row {i}, column {j}'
    doc.add_picture(str(images[n%8]),width=WordInches(3.5))
    doc.add_paragraph(f'END CAPACITY PAGE {n+1:02d}')
doc.save(out/'capacity-40-pages.docx')
prs=Presentation();prs.slide_width=Inches(13.333);prs.slide_height=Inches(7.5)
for n in range(20):
    slide=prs.slides.add_slide(prs.slide_layouts[6])
    title=slide.shapes.add_textbox(Inches(.6),Inches(.4),Inches(12),Inches(.7));title.text=f'SYNTHETIC CAPACITY SLIDE {n+1:02d}'
    title.text_frame.paragraphs[0].font.size=Pt(28)
    box=slide.shapes.add_textbox(Inches(.6),Inches(1.3),Inches(5),Inches(4));box.text='中文幻灯片容量验收\nSynthetic owned data only\nBullet 1: layout\nBullet 2: retained content'
    slide.shapes.add_picture(str(images[n%8]),Inches(6),Inches(1.6),width=Inches(6.4))
prs.save(out/'capacity-20-slides.pptx')
print(json.dumps({p.name:{'bytes':p.stat().st_size,'sha256':hashlib.sha256(p.read_bytes()).hexdigest()} for p in out.glob('capacity-*')},indent=2))
